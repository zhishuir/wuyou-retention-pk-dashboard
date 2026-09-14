from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "dist" / "data" / "report-data.json"
INBOX = ROOT / "待发布日报"
SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}


def newest_workbook() -> Path:
    candidates = (
        sorted(
            (path for path in INBOX.iterdir() if path.suffix.lower() in SUPPORTED_SUFFIXES),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        if INBOX.exists()
        else []
    )
    if not candidates:
        raise FileNotFoundError(f"请先把日报Excel放入：{INBOX}")
    return candidates[0]


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def integer(value: object) -> int:
    if value in (None, ""):
        return 0
    return int(float(value))


def parse_report_date(workbook_path: Path, override: str | None) -> str:
    if override:
        try:
            return datetime.strptime(override, "%Y-%m-%d").strftime("%Y-%m-%d")
        except ValueError as exc:
            raise ValueError("--date 必须使用 YYYY-MM-DD 格式。") from exc

    filename = workbook_path.stem
    patterns = (
        r"(?<!\d)(20\d{2})[年._-](\d{1,2})[月._-](\d{1,2})日?(?!\d)",
        r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)",
    )
    for pattern in patterns:
        match = re.search(pattern, filename)
        if match:
            year, month, day = (int(part) for part in match.groups())
            return date(year, month, day).strftime("%Y-%m-%d")

    chinese_date = re.search(r"(?<!\d)(\d{1,2})月(\d{1,2})日?(?!\d)", filename)
    if chinese_date:
        month, day = (int(part) for part in chinese_date.groups())
        return date(datetime.now().year, month, day).strftime("%Y-%m-%d")

    raise ValueError(
        "三列日报文件名中缺少日期。请命名为“2026-09-14三列日报.xlsx”，"
        "或运行时添加 --date 2026-09-14。"
    )


def excluded_names(payload: dict) -> set[str]:
    return {text(name) for name in payload.get("excluded", []) if text(name)}


def roster_from_payload(payload: dict) -> list[dict]:
    excluded = excluded_names(payload)
    registry: dict[str, dict] = {}
    for person in payload.get("roster", []):
        name = text(person.get("name"))
        if name and name not in excluded:
            registry[name] = {
                "name": name,
                "bigGroup": text(person.get("bigGroup")),
                "smallGroup": text(person.get("smallGroup")),
            }

    for day_data in payload.get("days", []):
        for person in day_data.get("personal", []):
            name = text(person.get("name"))
            if (
                name
                and name not in excluded
                and person.get("matched")
                and name not in registry
            ):
                registry[name] = {
                    "name": name,
                    "bigGroup": text(person.get("bigGroup")),
                    "smallGroup": text(person.get("smallGroup")),
                }
    return list(registry.values())


def read_detailed_day(workbook, roster: list[dict], excluded: set[str]) -> dict:
    sheet = workbook["计分明细"]
    imported: dict[str, dict] = {}
    report_date = ""
    for row in sheet.iter_rows(min_row=4, values_only=True):
        name = text(row[3])
        if not name or name in excluded:
            continue
        if not report_date:
            date_value = row[0]
            if isinstance(date_value, (date, datetime)):
                report_date = date_value.strftime("%Y-%m-%d")
            else:
                report_date = text(date_value)[:10]
        big_group = text(row[1]) or "待确认班组"
        small_group = text(row[2]) or "待确认小组"
        plus = integer(row[4])
        failure = integer(row[5])
        qc = integer(row[6])
        imported[name] = {
            "name": name,
            "bigGroup": big_group,
            "smallGroup": small_group,
            "plus": plus,
            "failure": failure,
            "qc": qc,
            "score": plus - failure - qc,
            "matched": big_group != "待确认班组",
        }

    if not report_date:
        raise ValueError("“计分明细”中没有可用日期。")
    if not imported:
        raise ValueError("“计分明细”中没有人员数据。")

    for person in roster:
        imported.setdefault(
            person["name"],
            {
                **person,
                "plus": 0,
                "failure": 0,
                "qc": 0,
                "score": 0,
                "matched": True,
            },
        )
    return {"date": report_date, "personal": list(imported.values())}


def compact_sheet(workbook):
    for sheet in workbook.worksheets:
        headers = [text(sheet.cell(row=1, column=column).value) for column in range(1, 4)]
        if (
            "成功" in headers[0]
            and ("失败" in headers[1] or "降档" in headers[1])
            and "质检" in headers[2]
        ):
            return sheet
    raise ValueError(
        "未识别到三列日报。第一行请依次填写：成功加分、失败扣分、质检扣分。"
    )


def names_in_cell(value: object) -> list[str]:
    if value is None:
        return []
    return [part.strip() for part in re.split(r"[\r\n]+", str(value)) if part.strip()]


def read_compact_day(
    workbook,
    workbook_path: Path,
    roster: list[dict],
    excluded: set[str],
    report_date_override: str | None,
) -> dict:
    if not roster:
        raise ValueError("网页数据中没有正式人员名单，无法匹配大组和小组。")

    sheet = compact_sheet(workbook)
    plus_names: Counter[str] = Counter()
    failure_names: Counter[str] = Counter()
    qc_names: Counter[str] = Counter()
    for row in sheet.iter_rows(min_row=2, max_col=3, values_only=True):
        plus_names.update(names_in_cell(row[0]))
        failure_names.update(names_in_cell(row[1]))
        qc_names.update(names_in_cell(row[2]))

    roster_by_name = {person["name"]: person for person in roster}
    event_names = (set(plus_names) | set(failure_names) | set(qc_names)) - excluded
    ordered_names = [person["name"] for person in roster]
    ordered_names.extend(sorted(event_names - set(ordered_names)))

    people: list[dict] = []
    for name in ordered_names:
        matched = name in roster_by_name
        membership = roster_by_name.get(
            name,
            {"bigGroup": "待确认班组", "smallGroup": "待确认小组"},
        )
        plus = plus_names[name]
        failure = failure_names[name]
        qc = qc_names[name]
        people.append(
            {
                "name": name,
                "bigGroup": membership["bigGroup"],
                "smallGroup": membership["smallGroup"],
                "plus": plus,
                "failure": failure,
                "qc": qc,
                "score": plus - failure - qc,
                "matched": matched,
            }
        )

    return {
        "date": parse_report_date(workbook_path, report_date_override),
        "personal": people,
    }


def read_day(
    workbook_path: Path,
    payload: dict,
    report_date_override: str | None = None,
) -> dict:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"找不到日报文件：{workbook_path}")
    if workbook_path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError("日报文件必须是 .xlsx 或 .xlsm 格式。")

    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    roster = roster_from_payload(payload)
    excluded = excluded_names(payload)
    if "计分明细" in workbook.sheetnames:
        return read_detailed_day(workbook, roster, excluded)
    return read_compact_day(
        workbook,
        workbook_path,
        roster,
        excluded,
        report_date_override,
    )


def update_roster(payload: dict, day_data: dict) -> list[dict]:
    roster = roster_from_payload(payload)
    known_names = {person["name"] for person in roster}
    for person in day_data["personal"]:
        if person.get("matched") and person["name"] not in known_names:
            roster.append(
                {
                    "name": person["name"],
                    "bigGroup": person["bigGroup"],
                    "smallGroup": person["smallGroup"],
                }
            )
            known_names.add(person["name"])
    return roster


def update_data(
    workbook_path: Path,
    report_date_override: str | None = None,
    data_file: Path = DATA_FILE,
) -> dict:
    if data_file.exists():
        payload = json.loads(data_file.read_text(encoding="utf-8"))
    else:
        payload = {"updatedAt": "", "roster": [], "days": []}

    day_data = read_day(workbook_path, payload, report_date_override)
    days = [item for item in payload.get("days", []) if item.get("date") != day_data["date"]]
    days.append(day_data)
    days.sort(key=lambda item: item["date"])
    roster = update_roster(payload, day_data)
    result = {
        "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "excluded": sorted(excluded_names(payload)),
        "roster": roster,
        "days": days,
    }
    data_file.parent.mkdir(parents=True, exist_ok=True)
    data_file.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    people = day_data["personal"]
    return {
        "date": day_data["date"],
        "people": len(people),
        "plus": sum(person["plus"] for person in people),
        "failure": sum(person["failure"] for person in people),
        "qc": sum(person["qc"] for person in people),
        "unmatched": [person["name"] for person in people if not person["matched"]],
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="把三列姓名日报或计分明细日报更新到网页数据文件。"
    )
    parser.add_argument(
        "workbook",
        nargs="?",
        type=Path,
        help="日报 Excel 路径；省略时读取“待发布日报”目录中最新的文件。",
    )
    parser.add_argument(
        "--date",
        help="三列日报对应日期，格式为 YYYY-MM-DD；省略时从文件名识别。",
    )
    parser.add_argument(
        "--data-file",
        type=Path,
        default=DATA_FILE,
        help=argparse.SUPPRESS,
    )
    args = parser.parse_args()
    try:
        workbook_path = args.workbook.resolve() if args.workbook else newest_workbook()
        summary = update_data(workbook_path, args.date, args.data_file.resolve())
        print(
            f"已更新 {summary['date']} 日报：成功加分 {summary['plus']} 次，"
            f"失败扣分 {summary['failure']} 次，质检扣分 {summary['qc']} 次，"
            f"名单共 {summary['people']} 人。"
        )
        if summary["unmatched"]:
            print("待确认名单：" + "、".join(summary["unmatched"]))
        return 0
    except Exception as exc:
        print(f"更新失败：{exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
