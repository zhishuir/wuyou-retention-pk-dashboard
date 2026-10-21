"""按两份独立分组表同步参赛名单，并重算已发布日期的人员归属。"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from update_pk_from_workbook import (
    PROJECT_SCORES,
    aggregate_personal,
    canonical_name as pk_name,
    parse_date_from_filename as pk_date,
    read_records,
)
from update_from_workbook import canonical_name as retention_name, read_day


ROOT = Path(__file__).resolve().parents[1]
REPORT_DATA = ROOT / "dist" / "data" / "report-data.json"
PK_DATA = ROOT / "dist" / "data" / "pk-data.json"
REPORT_ARCHIVE = ROOT / "待发布日报" / "已发布"
PK_ARCHIVE = ROOT / "待发布PK" / "已发布"
MEDIA_GROUPS = {"晶晶组", "左娜组"}
CONFIRMED_MARKETING_HEADCOUNT = {("扬帆组", "扬帆组1组"): 3}


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def validate_roster(roster: list[dict], title: str) -> None:
    duplicates = [name for name, count in Counter(p["name"] for p in roster).items() if count > 1]
    if duplicates:
        raise ValueError(f"{title}存在重名，请核实：{'、'.join(sorted(duplicates))}")
    if any(not p["bigGroup"] or not p["smallGroup"] for p in roster):
        raise ValueError(f"{title}存在未填写大组或小组的人员。")


def marketing_roster(path: Path) -> tuple[list[dict], list[str]]:
    sheet = load_workbook(path, data_only=True, read_only=True)["分组汇总表"]
    roster: list[dict] = []
    warnings: list[str] = []
    declared_total = None
    for big, small, declared, people in sheet.iter_rows(min_row=2, max_col=4, values_only=True):
        big = text(big)
        if big == "合计":
            declared_total = int(declared)
            break
        small = text(small)
        names = [pk_name(part) for part in re.split(r"[、，,；;\r\n]+", text(people)) if text(part)]
        if not big or not small or not names:
            raise ValueError(f"营销分组表存在空班组、小组或名单：{big} / {small}")
        if declared is not None and int(declared) != len(names):
            if CONFIRMED_MARKETING_HEADCOUNT.get((big, small)) != len(names):
                raise ValueError(f"营销分组表 {small} 标注 {declared} 人，实际 {len(names)} 人，请先核实。")
            warnings.append(f"{small}：表内标注 {declared} 人，按已确认的 {len(names)} 人计算")
        roster.extend({"name": name, "bigGroup": big, "smallGroup": small} for name in names)
    validate_roster(roster, "营销分组表")
    if declared_total is not None and declared_total != len(roster):
        if declared_total - len(roster) != len(warnings):
            raise ValueError(f"营销分组表合计 {declared_total} 人，实际 {len(roster)} 人，请先核实。")
        warnings.append(f"表内合计 {declared_total} 人，按已确认的 {len(roster)} 个姓名计算")
    return roster, warnings


def retention_roster(path: Path) -> list[dict]:
    sheet = load_workbook(path, data_only=True, read_only=True)["挽留王者PK赛分组表"]
    roster: list[dict] = []
    declared_counts: dict[tuple[str, str], int] = {}
    big = small = ""
    for row in sheet.iter_rows(min_row=2, max_col=4, values_only=True):
        if row[0] and not text(row[0]).startswith("AIGC:"):
            big = text(row[0])
        if row[1] and not text(row[1]).startswith("{"):
            small = text(row[1])
        name = text(row[2])
        if not name:
            continue
        if not big or not small:
            raise ValueError(f"挽留分组表中 {name} 缺少班组或小组。")
        if row[3] is not None:
            declared_counts[(big, small)] = int(row[3])
        roster.append({"name": retention_name(name), "bigGroup": big, "smallGroup": small})
    validate_roster(roster, "挽留分组表")
    actual_counts = Counter((p["bigGroup"], p["smallGroup"]) for p in roster)
    for group, count in actual_counts.items():
        if declared_counts.get(group) != count:
            raise ValueError(f"挽留分组表 {group[1]} 标注 {declared_counts.get(group)} 人，实际 {count} 人。")
    return roster


def report_day(day: dict, roster: list[dict], retired: set[str], excluded: set[str]) -> dict:
    known = {p["name"]: p for p in roster}
    existing: dict[str, dict] = {}
    pending: list[dict] = []
    for source in day.get("personal", []):
        name = retention_name(source.get("name"))
        if name in retired or name in excluded:
            continue
        if name not in known:
            item = deepcopy(source)
            item["name"] = name
            item["matched"] = False
            item["bigGroup"] = "待确认班组"
            item["smallGroup"] = "待确认小组"
            pending.append(item)
            continue
        item = existing.setdefault(name, {"plus": 0, "failure": 0, "qc": 0})
        for field in ("plus", "failure", "qc"):
            item[field] += int(source.get(field, 0))

    people = []
    for person in roster:
        values = existing.get(person["name"], {"plus": 0, "failure": 0, "qc": 0})
        people.append({
            **person,
            **values,
            "score": values["plus"] - values["failure"] - values["qc"],
            "matched": True,
        })
    return {**day, "personal": people + pending}


def pk_day(day: dict, roster: dict[str, dict], retired: set[str], excluded: set[str]) -> dict:
    people: dict[str, dict] = {}
    for source in day.get("personal", []):
        name = pk_name(source.get("name"))
        if name in retired or name in excluded:
            continue
        membership = roster.get(name)
        target = people.setdefault(name, {
            "name": name,
            "bigGroup": membership["bigGroup"] if membership else "待确认班组",
            "smallGroup": membership["smallGroup"] if membership else "待确认小组",
            "matched": membership is not None,
            "score": 0,
            "detail": {},
        })
        for project, count in (source.get("detail") or {}).items():
            target["detail"][project] = target["detail"].get(project, 0) + int(count)
        target["score"] += int(source.get("score", 0))
    ordered = sorted(people.values(), key=lambda p: (-p["score"], p["name"]))
    for index, person in enumerate(ordered, 1):
        person["rank"] = index
    return {**day, "personal": ordered}


def restore_new_people(
    report: dict,
    pk: dict,
    report_names: set[str],
    pk_names: set[str],
    report_archive: Path,
    pk_archive: Path,
) -> tuple[int, int, list[str]]:
    """仅从原始日报恢复旧程序排除、但现已正式参赛的人员记录。"""
    warnings: list[str] = []
    recovered_report = recovered_pk = 0
    report_days = {day["date"]: day for day in report["days"]}
    pk_days = {day["date"]: day for day in pk["days"]}
    roster = {person["name"]: person for person in pk["roster"]}
    report_seen: set[str] = set()
    pk_seen: set[str] = set()

    if report_names:
        for path in sorted(report_archive.glob("*.xlsx")):
            source_day = read_day(path, report)
            report_date = source_day["date"]
            if report_date not in report_days:
                continue
            if report_date in report_seen:
                raise ValueError(f"挽留赛归档目录有重复日期：{report_date}")
            report_seen.add(report_date)
            target = {person["name"]: person for person in report_days[report_date]["personal"]}
            for person in source_day["personal"]:
                if person["name"] not in report_names:
                    continue
                saved = target[person["name"]]
                for field in ("plus", "failure", "qc", "score"):
                    saved[field] = person[field]
                if saved["plus"] or saved["failure"] or saved["qc"]:
                    recovered_report += 1
        missing = sorted(set(report_days) - report_seen)
        if missing:
            warnings.append("挽留赛以下日期没有归档原表，新纳入人员只能维持原已发布数据：" + "、".join(missing))

    if pk_names:
        for path in sorted(pk_archive.glob("*.xlsx")):
            report_date = pk_date(path)
            if report_date not in pk_days:
                continue
            if report_date in pk_seen:
                raise ValueError(f"营销赛归档目录有重复日期：{report_date}")
            pk_seen.add(report_date)
            workbook = load_workbook(path, data_only=True, read_only=True)
            counts, _ = read_records(workbook)
            target = {person["name"]: person for person in pk_days[report_date]["personal"]}
            for name in pk_names & set(counts):
                detail = counts[name]
                score = sum(count * PROJECT_SCORES[project] for project, count in detail.items())
                target[name] = {
                    **roster[name],
                    "matched": True,
                    "score": score,
                    "detail": dict(sorted(detail.items())),
                }
                recovered_pk += 1
            ordered = sorted(target.values(), key=lambda p: (-p["score"], p["name"]))
            for index, person in enumerate(ordered, 1):
                person["rank"] = index
            pk_days[report_date]["personal"] = ordered
        missing = sorted(set(pk_days) - pk_seen)
        if missing:
            warnings.append("营销赛以下日期没有归档原表，新纳入人员只能维持原已发布数据：" + "、".join(missing))

    return recovered_report, recovered_pk, warnings


def synchronize(
    marketing_path: Path,
    retention_path: Path,
    report_path: Path,
    pk_path: Path,
    report_archive: Path = REPORT_ARCHIVE,
    pk_archive: Path = PK_ARCHIVE,
) -> tuple[dict, dict, list[str]]:
    marketing_calls, warnings = marketing_roster(marketing_path)
    retention_calls = retention_roster(retention_path)
    report = json.loads(report_path.read_text(encoding="utf-8"))
    pk = json.loads(pk_path.read_text(encoding="utf-8"))
    media = [deepcopy(p) for p in report["roster"] if p.get("bigGroup") in MEDIA_GROUPS]
    validate_roster(media, "新媒体名单")

    old_calls = {p["name"] for p in report["roster"] if p.get("bigGroup") not in MEDIA_GROUPS}
    previous_pk_roster = pk.get("roster", report["roster"])
    previous_pk_names = {p["name"] for p in previous_pk_roster}
    old_pk_calls = {p["name"] for p in previous_pk_roster if p.get("bigGroup") not in MEDIA_GROUPS}
    report_roster = retention_calls + media
    pk_roster = marketing_calls + deepcopy(media)
    validate_roster(report_roster, "挽留赛完整名单")
    validate_roster(pk_roster, "营销赛完整名单")

    report_retired = old_calls - {p["name"] for p in retention_calls}
    pk_retired = old_pk_calls - {p["name"] for p in marketing_calls}
    report_excluded = (set(report.get("excluded", [])) | report_retired) - {p["name"] for p in report_roster}
    pk_excluded = (set(pk.get("excluded", report.get("excluded", []))) | pk_retired) - {p["name"] for p in pk_roster}

    report["roster"] = report_roster
    report["excluded"] = sorted(report_excluded)
    report["days"] = [report_day(day, report_roster, report_retired, report_excluded) for day in report["days"]]
    report["updatedAt"] = datetime.now().strftime("%Y-%m-%d %H:%M")

    pk_roster_by_name = {p["name"]: p for p in pk_roster}
    pk["roster"] = pk_roster
    pk["excluded"] = sorted(pk_excluded)
    pk["days"] = [pk_day(day, pk_roster_by_name, pk_retired, pk_excluded) for day in pk["days"]]
    recovered_report, recovered_pk, recovery_warnings = restore_new_people(
        report,
        pk,
        {p["name"] for p in report_roster} - {p["name"] for p in media} - old_calls,
        {p["name"] for p in pk_roster} - previous_pk_names,
        report_archive,
        pk_archive,
    )
    warnings.extend(recovery_warnings)
    if recovered_report or recovered_pk:
        warnings.append(f"从归档原表恢复：挽留赛 {recovered_report} 人日、营销赛 {recovered_pk} 人日")
    pk["personal"] = aggregate_personal(pk["days"], pk_roster_by_name)
    pk["updatedAt"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    return report, pk, warnings


def main() -> int:
    parser = argparse.ArgumentParser(description="按独立分组表同步营销赛、挽留赛及历史排名。")
    parser.add_argument("marketing", type=Path, help="营销先锋PK赛分组表.xlsx")
    parser.add_argument("retention", type=Path, help="挽留王者PK赛分组表.xlsx")
    parser.add_argument("--check", action="store_true", help="仅核对名单，不写入数据文件")
    parser.add_argument("--report-data", type=Path, default=REPORT_DATA, help=argparse.SUPPRESS)
    parser.add_argument("--pk-data", type=Path, default=PK_DATA, help=argparse.SUPPRESS)
    args = parser.parse_args()
    try:
        report, pk, warnings = synchronize(args.marketing, args.retention, args.report_data, args.pk_data)
        for warning in warnings:
            print("[名单核对] " + warning)
        media_count = sum(p["bigGroup"] in MEDIA_GROUPS for p in report["roster"])
        print(f"挽留赛：{len(report['roster'])} 人；营销赛：{len(pk['roster'])} 人；新媒体保留 {media_count} 人。")
        print(f"历史日期：挽留赛 {len(report['days'])} 天；营销赛 {len(pk['days'])} 天。")
        if args.check:
            return 0
        args.report_data.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        args.pk_data.write_text(json.dumps(pk, ensure_ascii=False, indent=2), encoding="utf-8")
        print("两赛事名单与历史数据已同步。")
        return 0
    except Exception as exc:
        print(f"名单同步失败：{exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
