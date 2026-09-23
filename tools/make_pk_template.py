from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "营销PK赛模板.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="112B46")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)


def style_header(*cells) -> None:
    for cell in cells:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def main() -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "加分记录"
    ws.append(["姓名", "项目"])
    ws["A1"].comment = Comment("必填，员工姓名，须与挽留考核名单中的姓名一致。", "模板说明")
    ws["B1"].comment = Comment(
        "必填，项目名称，须为固定项目之一（分值由程序自动判定）。", "模板说明"
    )
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 36
    style_header(ws["A1"], ws["B1"])

    wb.save(OUTPUT)
    print(f"已生成模板：{OUTPUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
