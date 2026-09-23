from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PK_DATA_FILE = ROOT / "dist" / "data" / "pk-data.json"
RETENTION_DATA_FILE = ROOT / "dist" / "data" / "report-data.json"
INBOX = ROOT / "待发布PK"
SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}

NAME_HEADERS = ("姓名", "人名")
PROJECT_HEADERS = ("项目", "项目名称")

# 项目分值固定表：只加分不扣分，不同项目分值不同。
# 新增/调整项目时，在这里修改即可，模板里不需要再填分值。
PROJECT_SCORES: dict[str, int] = {
    "20元20G-7天提速流量包": 1,
    "30元30G-7天提速流量包": 1,
    "9.9元权益包": 1,
    "35元20G连续包月流量包（首月19.9元）": 2,
    "15元10G包月流量包": 2,
}


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def newest_workbook() -> Path:
    candidates = (
        sorted(
            (path for path in INBOX.iterdir() if path.suffix.lower() in SUPPORTED_SUFFIXES),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        if INBOX.exists()
        else []
    )
    if not candidates:
        raise FileNotFoundError(f"请先把营销PK赛模板放入：{INBOX}")
    return candidates[0]


def parse_date_from_filename(path: Path) -> str:
    patterns = (
        r"(?<!\d)(20\d{2})[年._-](\d{1,2})[月._-](\d{1,2})日?(?!\d)",
        r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)",
    )
    for pattern in patterns:
        match = re.search(pattern, path.stem)
        if match:
            year, month, day = (int(part) for part in match.groups())
            return date(year, month, day).isoformat()
    return ""


def read_records(workbook) -> tuple[dict, set[str]]:
    if "加分记录" not in workbook.sheetnames:
        raise ValueError("模板中缺少「加分记录」工作表。")
    sheet = workbook["加分记录"]

    header_row = next(sheet.iter_rows(min_row=1, max_row=1, max_col=3, values_only=True), None)
    headers = [text(cell) for cell in (header_row or [])]
    index = {"name": -1, "project": -1}
    for position, header in enumerate(headers):
        if header in NAME_HEADERS and index["name"] == -1:
            index["name"] = position
        elif header in PROJECT_HEADERS and index["project"] == -1:
            index["project"] = position
    if -1 in index.values():
        raise ValueError("表头需包含「姓名」和「项目」两列。")

    counts: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    unknown_projects: set[str] = set()
    for row in sheet.iter_rows(min_row=2, max_col=max(index.values()) + 1, values_only=True):
        name = text(row[index["name"]])
        project = text(row[index["project"]])
        if not name:
            continue
        if not project:
            continue
        if project not in PROJECT_SCORES:
            unknown_projects.add(project)
            continue
        counts[name][project] += 1
    return dict(counts), unknown_projects


def roster_from_retention() -> dict[str, dict]:
    payload = json.loads(RETENTION_DATA_FILE.read_text(encoding="utf-8"))
    excluded = {text(name) for name in payload.get("excluded", []) if text(name)}
    roster: dict[str, dict] = {}
    for person in payload.get("roster", []):
        name = text(person.get("name"))
        if name and name not in excluded:
            roster[name] = {
                "name": name,
                "bigGroup": text(person.get("bigGroup")) or "待确认班组",
                "smallGroup": text(person.get("smallGroup")) or "待确认小组",
            }
    return roster


def build_personal(counts: dict, roster: dict) -> list[dict]:
    people: list[dict] = []
    for name, detail in counts.items():
        total = sum(count * PROJECT_SCORES[project] for project, count in detail.items())
        membership = roster.get(name)
        people.append(
            {
                "name": name,
                "bigGroup": membership["bigGroup"] if membership else "待确认班组",
                "smallGroup": membership["smallGroup"] if membership else "待确认小组",
                "matched": membership is not None,
                "score": total,
                "detail": dict(sorted(detail.items())),
            }
        )
    people.sort(key=lambda person: (-person["score"], person["name"]))
    for index, person in enumerate(people):
        person["rank"] = index + 1
    return people


def build_project_scores() -> list[dict]:
    return [
        {"name": name, "score": score}
        for name, score in sorted(PROJECT_SCORES.items(), key=lambda item: (-item[1], item[0]))
    ]


def update_data(workbook_path: Path, data_file: Path = PK_DATA_FILE) -> dict:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    counts, unknown_projects = read_records(workbook)
    if unknown_projects:
        print("警告：以下项目不在固定分值表中，已忽略：" + "、".join(sorted(unknown_projects)))
    roster = roster_from_retention()
    if not counts:
        raise ValueError("「加分记录」表中没有有效数据。")
    personal = build_personal(counts, roster)
    payload = {
        "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "date": parse_date_from_filename(workbook_path),
        "projectScores": build_project_scores(),
        "personal": personal,
    }
    data_file.parent.mkdir(parents=True, exist_ok=True)
    data_file.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    unmatched = [person["name"] for person in personal if not person["matched"]]
    return {
        "date": payload["date"],
        "people": len(personal),
        "total": sum(person["score"] for person in personal),
        "unmatched": unmatched,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="把营销PK赛模板(姓名+项目)更新到网页数据文件。")
    parser.add_argument(
        "workbook",
        nargs="?",
        type=Path,
        help="模板 Excel 路径；省略时读取“待发布PK”目录中最新的文件。",
    )
    parser.add_argument("--data-file", type=Path, default=PK_DATA_FILE, help=argparse.SUPPRESS)
    args = parser.parse_args()
    try:
        workbook_path = args.workbook.resolve() if args.workbook else newest_workbook()
        summary = update_data(workbook_path, args.data_file.resolve())
        day_text = f"（{summary['date']}）" if summary["date"] else ""
        print(
            f"已更新营销PK赛{day_text}：{summary['people']} 人参赛，累计加分 {summary['total']} 分。"
        )
        if summary["unmatched"]:
            print("待确认名单：" + "、".join(summary["unmatched"]))
        return 0
    except Exception as exc:
        print(f"PK赛更新失败：{exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
